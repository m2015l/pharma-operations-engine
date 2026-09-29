"""Exporta a conversa (ou uma resposta) para Word — frontend/PLANO_V2.md §4.

Gerador puro: `gerar(ExportarDocxIn) -> bytes`. Só formata o que o navegador
mandou (o mesmo `RespostaOut` que /perguntar devolveu); não chama LLM nem
banco, então custa zero também no modo real.

Layout: A4, margens de 2,5 cm, Lato 11 pt (o Word troca por Calibri se não
tiver a fonte). Cabeçalho com o logo e "Licitações" sobre uma linha #124443;
rodapé com o aviso de IA e "p. X/Y". Cada turno: pergunta, resposta (markdown
mínimo: parágrafo, lista, **negrito**, [n]), alertas, régua em texto e a
tabela das fontes citadas com o trecho (até ~400 caracteres). Sem autor.
"""

from __future__ import annotations

import io
import re
from datetime import UTC, datetime, timedelta, timezone
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from schemas import CitacaoOut, ExportarDocxIn, RespostaOut, TurnoExportado

LOGO = Path(__file__).with_name("assets") / "logo.png"

PETROLEO = RGBColor(0x12, 0x44, 0x43)
TEAL_TEXTO = RGBColor(0x27, 0x71, 0x60)
TINTA = RGBColor(0x21, 0x25, 0x29)
GRAFITE = RGBColor(0x5B, 0x6B, 0x6E)
ALERTA = RGBColor(0x7A, 0x5A, 0x0A)
ERRO = RGBColor(0xB4, 0x23, 0x18)
FONTE = "Lato"

AVISO = (
    "Resposta gerada por IA a partir do acervo da Solufarma. "
    "Confira sempre no documento original antes de usar."
)
MAX_TRECHO = 400


def _fuso_sp() -> timezone:
    # zoneinfo no Windows depende do pacote tzdata; o Brasil não tem horário
    # de verão desde 2019, então -03:00 fixo é exato.
    try:
        from zoneinfo import ZoneInfo

        return ZoneInfo("America/Sao_Paulo")  # type: ignore[return-value]
    except (ImportError, LookupError):  # ZoneInfoNotFoundError é um KeyError
        return timezone(timedelta(hours=-3), "BRT")


def _local(dt: datetime) -> datetime:
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=UTC)
    return dt.astimezone(_fuso_sp())


def nome_arquivo(corpo: ExportarDocxIn) -> str:
    prefixo = "resposta" if corpo.escopo == "resposta" else "consulta"
    return f"{prefixo}-solufarma-{_local(corpo.gerado_em):%Y-%m-%d-%H%M}.docx"


# --------------------------------------------------------------------------
# Texto
# --------------------------------------------------------------------------


def limpar_trecho(trecho: str, limite: int = MAX_TRECHO) -> str:
    """Mesma limpeza do `limparTrecho` do painel de fontes + corte em ~400 caracteres."""
    t = trecho.replace("**", "")
    t = re.sub(r"^#+\s*", "", t, flags=re.MULTILINE)
    t = re.sub(r"\n\s*\n+", "\n", t).strip()
    if len(t) <= limite:
        return t
    corte = t.rfind(" ", 0, limite)
    return t[: corte if corte > limite // 2 else limite].rstrip(" ,;.:-") + "…"


def paginas(c: CitacaoOut) -> str:
    if c.pagina_fim != c.pagina_inicio:
        return f"p. {c.pagina_inicio}–{c.pagina_fim}"
    return f"p. {c.pagina_inicio}"


_ROTULO_TIPO = {
    "atestado": "Atestado / CAT",
    "edital": "Edital",
    "ata": "Ata",
    "resultado": "Resultado",
    "anexo": "Anexo",
    "proposta": "Proposta",
    "recurso": "Recurso",
    "outro": "Outro",
}

_INLINE = re.compile(r"(\*\*.+?\*\*|\[\d+\])")
_ITEM = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")
_TITULO_MD = re.compile(r"^#{1,6}\s+(.*)$")


def _fonte_run(run, tamanho: float | None = None, cor: RGBColor | None = None) -> None:
    run.font.name = FONTE
    if tamanho:
        run.font.size = Pt(tamanho)
    if cor is not None:
        run.font.color.rgb = cor


def _inline(par, texto: str, validos: set[int], negrito: bool = False) -> None:
    """**negrito** e marcadores [n]; marcador sem fonte sai riscado + '(fonte inexistente)'."""
    for parte in _INLINE.split(texto):
        if not parte:
            continue
        if parte.startswith("**") and parte.endswith("**") and len(parte) > 4:
            _inline(par, parte[2:-2], validos, negrito=True)
            continue
        m = re.fullmatch(r"\[(\d+)\]", parte)
        if m:
            n = int(m.group(1))
            run = par.add_run(f"[{n}]")
            _fonte_run(run, cor=TEAL_TEXTO if n in validos else ALERTA)
            run.bold = True
            if n not in validos:
                run.font.strike = True
                aviso = par.add_run(" (fonte inexistente)")
                _fonte_run(aviso, 9, ALERTA)
            continue
        run = par.add_run(parte)
        _fonte_run(run)
        run.bold = negrito or None


def _markdown(doc, texto: str, validos: set[int]) -> None:
    """Conversor mínimo: parágrafo, título #, lista -/1., **negrito** e [n]."""
    for bloco in re.split(r"\n\s*\n", texto.strip()):
        for linha in bloco.split("\n"):
            if not linha.strip():
                continue
            if m := _TITULO_MD.match(linha):
                par = doc.add_paragraph()
                _inline(par, m.group(1), validos, negrito=True)
                continue
            if m := _ITEM.match(linha):
                recuo, marca, conteudo = m.groups()
                nivel = 1 + len(recuo.replace("\t", "  ")) // 2
                par = doc.add_paragraph()
                par.paragraph_format.left_indent = Cm(0.6 * nivel + 0.2)
                par.paragraph_format.first_line_indent = Cm(-0.5)
                par.paragraph_format.space_after = Pt(3)
                simbolo = "•" if marca in "-*+" else marca.rstrip(")").rstrip(".") + "."
                _fonte_run(par.add_run(f"{simbolo}\t"))
                par.paragraph_format.tab_stops.add_tab_stop(Cm(0.6 * nivel + 0.2))
                _inline(par, conteudo, validos)
                continue
            par = doc.add_paragraph()
            _inline(par, linha.strip(), validos)


# --------------------------------------------------------------------------
# Peças de XML que o python-docx não expõe
# --------------------------------------------------------------------------


def _campo(par, instrucao: str, tamanho: float, cor: RGBColor) -> None:
    """Campo PAGE/NUMPAGES (o Word calcula ao abrir e ao imprimir)."""

    def run_com(elemento) -> None:
        run = par.add_run()
        _fonte_run(run, tamanho, cor)
        run._r.append(elemento)

    def fld_char(tipo: str):
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), tipo)
        return el

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instrucao} "
    provisorio = OxmlElement("w:t")
    provisorio.text = "1"

    run_com(fld_char("begin"))
    run_com(instr)
    run_com(fld_char("separate"))
    run_com(provisorio)
    run_com(fld_char("end"))


def _borda_inferior(par, cor: str = "124443", espessura: int = 12) -> None:
    ppr = par._p.get_or_add_pPr()
    bordas = OxmlElement("w:pBdr")
    baixo = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", str(espessura)), ("space", "4"), ("color", cor)):
        baixo.set(qn(f"w:{k}"), v)
    bordas.append(baixo)
    ppr.append(bordas)


def _sombrear(celula, cor: str) -> None:
    tcpr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    tcpr.append(shd)


def _linha_cabecalho(linha) -> None:
    """Repete no topo de cada página e não se separa da primeira fonte."""
    trpr = linha._tr.get_or_add_trPr()
    for tag in ("w:tblHeader", "w:cantSplit"):
        trpr.append(OxmlElement(tag))


def _nao_quebrar(linha) -> None:
    linha._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def _bordas_tabela(tabela) -> None:
    """Só linhas horizontais (sem bordas verticais), como a tabela do Contratos."""
    tblpr = tabela._tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        if lado in ("left", "right", "insideV"):
            el.set(qn("w:val"), "nil")
        else:
            for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "DEE2E6")):
                el.set(qn(f"w:{k}"), v)
        bordas.append(el)
    tblpr.append(bordas)


# --------------------------------------------------------------------------
# Documento
# --------------------------------------------------------------------------


def _preparar(doc) -> None:
    secao = doc.sections[0]
    secao.page_width, secao.page_height = Cm(21), Cm(29.7)
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(secao, lado, Cm(2.5))
    secao.header_distance = Cm(1.2)
    secao.footer_distance = Cm(1.2)

    normal = doc.styles["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(11)
    normal.font.color.rgb = TINTA
    rfonts = normal.element.rPr.get_or_add_rFonts()
    for atributo in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(atributo), FONTE)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    largura = secao.page_width - secao.left_margin - secao.right_margin

    cab = secao.header.paragraphs[0]
    if LOGO.is_file():
        cab.add_run().add_picture(str(LOGO), height=Cm(0.75))
        separador = cab.add_run("   |   Licitações")
    else:
        separador = cab.add_run("SOLUFARMA ENGENHARIA · Licitações")
    _fonte_run(separador, 12, PETROLEO)
    separador.bold = True
    _borda_inferior(cab)

    rod = secao.footer.paragraphs[0]
    rod.paragraph_format.tab_stops.add_tab_stop(largura, WD_TAB_ALIGNMENT.RIGHT)
    _fonte_run(rod.add_run(AVISO + "\t"), 8, GRAFITE)
    _fonte_run(rod.add_run("p. "), 8, GRAFITE)
    _campo(rod, "PAGE", 8, GRAFITE)
    _fonte_run(rod.add_run("/"), 8, GRAFITE)
    _campo(rod, "NUMPAGES", 8, GRAFITE)


def _tabela_fontes(doc, citacoes: list[CitacaoOut]) -> None:
    rotulo = doc.add_paragraph()
    rotulo.paragraph_format.space_before = Pt(8)
    rotulo.paragraph_format.keep_with_next = True
    run = rotulo.add_run("Fontes")
    _fonte_run(run, 11, PETROLEO)
    run.bold = True

    tabela = doc.add_table(rows=1, cols=2)
    tabela.alignment = WD_TABLE_ALIGNMENT.LEFT
    tabela.autofit = False
    _bordas_tabela(tabela)
    larguras = (Cm(1.2), Cm(14.8))
    _linha_cabecalho(tabela.rows[0])
    for i, texto in enumerate(("Nº", "Documento, página e trecho citado")):
        celula = tabela.rows[0].cells[i]
        celula.width = larguras[i]
        _sombrear(celula, "124443")
        celula.paragraphs[0].paragraph_format.keep_with_next = True
        run = celula.paragraphs[0].add_run(texto)
        _fonte_run(run, 9, RGBColor(0xFF, 0xFF, 0xFF))
        run.bold = True

    for c in citacoes:
        _nao_quebrar(tabela.add_row())
        linha = tabela.rows[-1].cells
        linha[0].width, linha[1].width = larguras
        marcador = linha[0].paragraphs[0].add_run(f"[{c.marcador}]")
        _fonte_run(marcador, 10, TEAL_TEXTO)
        marcador.bold = True

        cab = linha[1].paragraphs[0]
        cab.paragraph_format.space_after = Pt(2)
        orgao = cab.add_run(f"{c.orgao} — ")
        _fonte_run(orgao, 10)
        orgao.bold = True
        _fonte_run(cab.add_run(c.titulo_documento), 10)

        meta = linha[1].add_paragraph()
        meta.paragraph_format.space_after = Pt(2)
        partes = [_ROTULO_TIPO.get(c.tipo_documento, c.tipo_documento)]
        partes.append(c.secao or "folha de rosto (sem seção)")
        partes.append(paginas(c))
        _fonte_run(meta.add_run(" · ".join(partes)), 9, GRAFITE)

        trecho = limpar_trecho(c.trecho)
        if trecho:
            par = linha[1].add_paragraph()
            par.paragraph_format.left_indent = Cm(0.3)
            run = par.add_run(f"“{trecho.replace(chr(10), ' ')}”")
            _fonte_run(run, 9, TINTA)
            run.italic = True


def _resposta(doc, r: RespostaOut) -> None:
    rotulo = doc.add_paragraph()
    rotulo.paragraph_format.keep_with_next = True
    run = rotulo.add_run("Resposta")
    _fonte_run(run, 10, GRAFITE)
    run.bold = True

    if not r.encontrou_no_acervo:
        par = doc.add_paragraph()
        run = par.add_run("Nada encontrado no acervo")
        _fonte_run(run, 10, TINTA)
        run.bold = True

    validos = {c.marcador for c in r.citacoes}
    _markdown(doc, r.texto, validos)

    if r.encontrou_no_acervo and not r.contrato_citacao_ok:
        par = doc.add_paragraph()
        detalhe = (
            f"O agente citou trechos que não existem ({', '.join(map(str, r.citacoes_invalidas))})."
            if r.citacoes_invalidas
            else "O agente não citou nenhum trecho."
        )
        run = par.add_run(f"⚠ Resposta sem fonte rastreável. {detalhe} Não use sem conferir no documento.")
        _fonte_run(run, 10, ALERTA)
        run.bold = True

    if r.encontrou_no_acervo:
        par = doc.add_paragraph()
        n, m = len(r.citacoes), r.trechos_recuperados
        texto = (
            f"{n} {'citada' if n == 1 else 'citadas'} de {m} "
            f"{'trecho recuperado' if m == 1 else 'trechos recuperados'}"
        )
        _fonte_run(par.add_run(texto), 9, GRAFITE)
        if r.citacoes:
            _tabela_fontes(doc, r.citacoes)


def _turno(doc, numero: int, turno: TurnoExportado, singular: bool) -> None:
    titulo = doc.add_paragraph()
    titulo.paragraph_format.space_before = Pt(14)
    titulo.paragraph_format.keep_with_next = True
    run = titulo.add_run("Pergunta" if singular else f"Pergunta {numero}")
    _fonte_run(run, 12, PETROLEO)
    run.bold = True

    pergunta = doc.add_paragraph()
    pergunta.paragraph_format.keep_with_next = True
    run = pergunta.add_run(turno.pergunta)
    _fonte_run(run)
    run.bold = True

    if turno.resposta is not None:
        _resposta(doc, turno.resposta)
    else:
        par = doc.add_paragraph()
        motivo = turno.erro or "a consulta não terminou"
        _fonte_run(par.add_run(f"Sem resposta (erro: {motivo})"), 10, ERRO)


def gerar(corpo: ExportarDocxIn) -> bytes:
    doc = Document()
    _preparar(doc)
    singular = corpo.escopo == "resposta"
    quando = _local(corpo.gerado_em)

    # Sem autor (Roger, 29/09): o template do python-docx grava "python-docx".
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.comments = ""
    props.title = "Resposta do acervo técnico" if singular else "Consulta ao acervo técnico"
    props.subject = "Licitações · Solufarma"
    props.created = props.modified = quando

    titulo = doc.add_paragraph()
    titulo.paragraph_format.space_after = Pt(2)
    run = titulo.add_run("Resposta do acervo técnico" if singular else "Consulta ao acervo técnico")
    _fonte_run(run, 20, TINTA)

    n = len(corpo.turnos)
    meta = doc.add_paragraph()
    texto_meta = f"Gerado em {quando:%d/%m/%Y} às {quando:%H:%M}"
    if not singular:
        texto_meta += f" · {n} {'pergunta' if n == 1 else 'perguntas'}"
    _fonte_run(meta.add_run(texto_meta), 10, GRAFITE)

    if any(t.resposta and t.resposta.modo == "demo" for t in corpo.turnos):
        demo = doc.add_paragraph()
        run = demo.add_run("Modo demonstração — respostas gravadas das auditorias")
        _fonte_run(run, 9, ALERTA)
        run.bold = True

    for i, turno in enumerate(corpo.turnos, start=1):
        if i > 1:
            _borda_inferior(doc.add_paragraph(), cor="DEE2E6", espessura=6)
        _turno(doc, i, turno, singular)

    saida = io.BytesIO()
    doc.save(saida)
    return saida.getvalue()
